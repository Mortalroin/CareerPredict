from __future__ import annotations

import re
from collections import OrderedDict
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.schemas.job import JobImportItem


class JobSpreadsheetService:
    def parse_excel(self, file_bytes: bytes) -> list[JobImportItem]:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [self._normalize_header(value) for value in rows[0]]
        items: list[JobImportItem] = []

        for row in rows[1:]:
            if not any(row):
                continue

            row_map = self._row_to_map(headers, row)
            title = self._pick(row_map, ["job_title", "title", "position_name"])
            company_name = self._pick(row_map, ["company_name", "company"])
            if not title or not company_name:
                continue

            city = self._extract_city(self._pick(row_map, ["city", "location", "work_city"]))
            salary_text = self._pick(row_map, ["salary", "salary_range"])
            salary_min, salary_max = self._parse_salary_range(salary_text)
            description = self._pick(row_map, ["job_detail", "job_description", "description"])
            company_detail = self._pick(row_map, ["company_detail", "company_description"])
            full_text = f"{title}\n{description}\n{company_detail}".strip()

            degree_raw = self._pick(row_map, ["degree", "education_requirement"]) or self._extract_degree_text(full_text)
            required_skills = self._extract_skills_from_text(full_text)
            preferred_skills = self._extract_preferred_skills(company_detail)
            required_languages = self._extract_languages(full_text)
            min_years_experience = self._extract_years_experience(full_text)

            items.append(
                JobImportItem(
                    title=title,
                    company_name=company_name,
                    city=city,
                    industry=self._extract_industry_tags(full_text),
                    job_type="full_time",
                    salary_min=salary_min,
                    salary_max=salary_max,
                    degree_name_raw=degree_raw,
                    required_languages=required_languages,
                    required_skills=required_skills[:12],
                    preferred_skills=preferred_skills[:8],
                    required_tools=[
                        skill for skill in required_skills if skill in {"word", "excel", "figma", "photoshop", "cad"}
                    ],
                    responsibilities=self._extract_responsibilities(description),
                    description=full_text,
                )
            )

        return items

    @staticmethod
    def _row_to_map(headers: list[str], row: tuple[Any, ...]) -> dict[str, str]:
        mapped: dict[str, str] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            value = ""
            if index < len(row) and row[index] is not None:
                value = str(row[index]).strip()
            mapped[header] = value
        return mapped

    @staticmethod
    def _pick(row_map: dict[str, str], keys: list[str]) -> str:
        for key in keys:
            value = row_map.get(key, "")
            if value:
                return value
        return ""

    @staticmethod
    def _normalize_header(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        alias_map = {
            "\u5c97\u4f4d\u540d\u79f0": "job_title",
            "\u804c\u4f4d\u540d\u79f0": "job_title",
            "\u5c97\u4f4d": "job_title",
            "\u516c\u53f8\u540d\u79f0": "company_name",
            "\u516c\u53f8": "company_name",
            "\u5de5\u4f5c\u5730\u70b9": "city",
            "\u5730\u70b9": "city",
            "\u5730\u5740": "city",
            "\u533a\u57df": "city",
            "\u5de5\u4f5c\u57ce\u5e02": "city",
            "\u85aa\u8d44\u8303\u56f4": "salary",
            "\u85aa\u8d44": "salary",
            "\u5b66\u5386": "degree",
            "\u5b66\u5386\u8981\u6c42": "degree",
            "\u5c97\u4f4d\u8be6\u60c5": "job_detail",
            "\u804c\u4f4d\u8be6\u60c5": "job_detail",
            "\u804c\u4f4d\u63cf\u8ff0": "job_detail",
            "\u516c\u53f8\u8be6\u60c5": "company_detail",
            "\u516c\u53f8\u4ecb\u7ecd": "company_detail",
            "\u5c97\u4f4d\u8be6\u60c5\u5730\u5740": "job_url",
            "\u804c\u4f4d\u94fe\u63a5": "job_url",
            "\u66f4\u65b0\u65f6\u95f4": "updated_at",
        }
        return alias_map.get(text, text.strip().lower().replace(" ", "_"))

    @staticmethod
    def _extract_city(value: str) -> str:
        if not value:
            return ""
        parts = value.replace("/", " ").replace("-", " ").split()
        return parts[0] if parts else value

    @staticmethod
    def _parse_salary_range(value: str) -> tuple[int | None, int | None]:
        if not value:
            return None, None
        text = value.lower().replace("k", "000")
        numbers = []
        current = ""
        for char in text:
            if char.isdigit():
                current += char
            else:
                if current:
                    numbers.append(int(current))
                    current = ""
        if current:
            numbers.append(int(current))
        if len(numbers) >= 2:
            return numbers[0], numbers[1]
        if len(numbers) == 1:
            return numbers[0], numbers[0]
        return None, None

    @staticmethod
    def _extract_degree_text(text: str) -> str:
        patterns = [
            r"((?:\u7edf\u62db)?(?:\u672c\u79d1|\u7855\u58eb|\u535a\u58eb|\u5927\u4e13|\u4e13\u79d1|\u9ad8\u4e2d).{0,12}(?:\u53ca\u4ee5\u4e0a|\u4ee5\u4e0a|\u4f18\u5148)?)",
            r"((?:bachelor|master|phd).{0,20}(?:above|preferred)?)",
        ]
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return match.group(1).strip()
        return ""

    @staticmethod
    def _extract_years_experience(text: str) -> float:
        if not text:
            return 0.0
        patterns = [
            r"(\d+(?:\.\d+)?)\+?\s*\u5e74",
            r"\u6bd5\u4e1a\u6ee1\s*(\d+(?:\.\d+)?)\s*\u5e74",
            r"(\d+(?:\.\d+)?)\+?\s*(?:years|year)",
        ]
        lower = text.lower()
        values: list[float] = []
        for pattern in patterns:
            for match in re.findall(pattern, lower, re.IGNORECASE):
                try:
                    values.append(float(match))
                except ValueError:
                    continue
        return max(values) if values else 0.0

    @staticmethod
    def _extract_languages(text: str) -> list[str]:
        lower = text.lower()
        mapping = {
            "english": ["\u82f1\u8bed", "english", "cet-4", "cet-6"],
            "japanese": ["\u65e5\u8bed", "japanese"],
            "mandarin": ["\u666e\u901a\u8bdd", "mandarin"],
            "cantonese": ["\u7ca4\u8bed", "cantonese"],
        }
        return [label for label, keywords in mapping.items() if any(keyword in lower for keyword in keywords)]

    @staticmethod
    def _extract_industry_tags(text: str) -> list[str]:
        lower = text.lower()
        tags = []
        rules = {
            "internet": ["\u4e92\u8054\u7f51", "internet", "\u8f6f\u4ef6"],
            "manufacturing": ["\u5236\u9020", "\u751f\u4ea7", "manufacturing"],
            "foreign_trade": ["\u5916\u8d38"],
            "ai": ["ai", "\u4eba\u5de5\u667a\u80fd", "llm", "\u5927\u6a21\u578b"],
            "education": ["\u6559\u80b2"],
            "finance": ["\u91d1\u878d", "\u94f6\u884c"],
            "testing": ["\u6d4b\u8bd5", "test"],
        }
        for tag, keywords in rules.items():
            if any(keyword in lower for keyword in keywords):
                tags.append(tag)
        return tags

    @staticmethod
    def _extract_skills_from_text(text: str) -> list[str]:
        lower = text.lower()
        skills = [
            "python",
            "java",
            "sql",
            "excel",
            "word",
            "photoshop",
            "cad",
            "figma",
            "fastapi",
            "react",
            "vue",
            "\u5ba2\u6237\u670d\u52a1",
            "\u8ba2\u5355\u8ddf\u5355",
            "\u62a5\u5173",
            "\u51fa\u8d27\u5b89\u6392",
            "\u56de\u6b3e",
            "\u6295\u8bc9\u5904\u7406",
            "\u5916\u8d38",
            "\u8f6f\u4ef6\u6d4b\u8bd5",
            "\u529f\u80fd\u6d4b\u8bd5",
            "\u63a5\u53e3\u6d4b\u8bd5",
            "\u6570\u636e\u6d4b\u8bd5",
            "\u81ea\u52a8\u5316\u6d4b\u8bd5",
            "product planning",
            "data analysis",
            "rag",
            "langchain",
            "chromadb",
        ]
        found = [skill for skill in skills if skill in lower]
        return list(OrderedDict.fromkeys(found))

    @staticmethod
    def _extract_preferred_skills(text: str) -> list[str]:
        lower = text.lower()
        preferred = [
            "product planning",
            "data analysis",
            "sql",
            "python",
            "excel",
            "\u8ba1\u7b97\u673a\u4e13\u4e1a",
            "\u8f6f\u4ef6\u6d4b\u8bd5",
            "\u529f\u80fd\u6d4b\u8bd5",
        ]
        found = [skill for skill in preferred if skill in lower]
        return list(OrderedDict.fromkeys(found))

    @staticmethod
    def _extract_responsibilities(text: str) -> list[str]:
        if not text:
            return []
        lines = [line.strip(" ;；。.\t") for line in text.splitlines() if line.strip()]
        numbered = [line for line in lines if line[:1].isdigit()]
        if numbered:
            return numbered[:8]
        return lines[:5]
